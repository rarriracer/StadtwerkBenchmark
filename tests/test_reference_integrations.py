import io
import json
import zipfile
from datetime import date, datetime, timedelta, timezone
import pytest
from pptx import Presentation
from stadtwerke import reference_sources as sources, references
from stadtwerke.calculation import calculate_reference, available_metrics
from stadtwerke.reporting import add_topic


def record(payload):
    return {'text':json.dumps(payload) if not isinstance(payload,str) else payload,'url':'https://example.org/data',
        'fetched_at':'2026-10-07T12:00:00Z','sha256':'abc'}


def months():
    return [{'date':f'{y}-{m:02}-01','value':100} for y in (2024,2025) for m in range(1,13)]+[
        {'date':f'2026-{m:02}-01','value':110} for m in range(1,10)]


def test_flow_annual_anchor_and_growth_use_different_denominators():
    pair=sources.monthly_pair(months(),2024,2026,date(2026,10,7))
    ref=dict(pair,baseline=pair['baseline_annual'],aligned_baseline=pair['baseline'],kind='flow',unit='MWh')
    calc=calculate_reference({'value':120,'unit':'MWh'},ref)
    assert calc['expected']==pytest.approx(99)
    assert calc['historical_company_aligned']==90
    assert calc['company_change']==pytest.approx(9)
    assert calc['reference_change_percent']==pytest.approx(10)


def test_missing_month_stops_period_and_incomplete_base_is_rejected():
    rows=[r for r in months() if r['date']!='2026-03-01']
    assert sources.monthly_pair(rows,2025,2026,date(2026,10,7))['months']==2
    with pytest.raises(ValueError,match='vollständig'):
        sources.monthly_pair([r for r in rows if r['date']!='2025-12-01'],2025,2026,date(2026,10,7))


def test_context_never_scales_company_and_rate_changes_are_percentage_points():
    calc=calculate_reference({'value':5_000_000,'unit':'EUR'},{'kind':'context','baseline':3.17,'target':3.28})
    assert calc['expected'] is None and calc['ratio'] is None
    assert calc['reference_change']==pytest.approx(.11)
    with pytest.raises(ValueError,match='Einheiten'):
        calculate_reference({'value':5,'unit':'EUR'},{'kind':'flow','unit':'MWh','baseline':1,'target':2})


def test_gas_daily_average_conversion_includes_leap_day(monkeypatch):
    monkeypatch.setattr(sources,'get_record',lambda *a,**k:record('Monat;2024;2025;2026\n2;1,5;1,5;1,5\n'))
    rows,_=sources.gas_monthly('slp')
    assert [r['value'] for r in rows]==[43500,42000,42000]


def test_kba_requires_all_states_and_honors_confidentiality_and_pagination(monkeypatch):
    features=[{'attributes':{'Berichtszeitpunkt':'2026.07','Bundesland':s,'Pkw_BEV':100,
        'ZS_Pkw_BEV':'','ZS_Pkw_Plug_In_Hybrid':''}} for s in sorted(sources.GERMAN_STATES)]
    def fetch(url,params):
        offset=params['resultOffset'];return record({'features':features[offset:offset+8],'exceededTransferLimit':offset==0})
    monkeypatch.setattr(sources,'get_record',fetch)
    rows,_=sources.kba_rows();assert rows[0]['value']==1600
    features[0]['attributes']['ZS_Pkw_BEV']='.';features[0]['attributes']['Pkw_BEV']=0
    with pytest.raises(ValueError,match='freigegebener'): sources.kba_rows()


def test_spot_duration_weighting_handles_dst_and_resolution_change(monkeypatch):
    start=datetime(2025,3,1,tzinfo=sources.BERLIN).astimezone(timezone.utc)
    end=datetime(2025,4,1,tzinfo=sources.BERLIN).astimezone(timezone.utc)
    rows=[];instant=start
    while instant<end:
        step=timedelta(hours=1) if instant<start+timedelta(hours=100) else timedelta(minutes=15)
        rows.append({'timestamp':instant.isoformat(),'values':{'day_ahead_price':10 if step.total_seconds()==3600 else 20}})
        instant+=step
    monkeypatch.setattr(sources,'get_record',lambda *a,**k:record({'endpoint':'price','unit':'EUR/MWh','data':rows}))
    monthly,_=sources.spot_monthly(2025,2025,date(2025,4,7))
    assert monthly[0]['hours']==743
    assert monthly[0]['value']==pytest.approx((100*10+643*20)/743)
    rows[100]['values']['day_ahead_price']=None
    with pytest.raises(ValueError,match='Fehlender'): sources.spot_monthly(2025,2025,date(2025,4,7))


def test_public_charger_exact_city_state_and_national_fallback():
    # A small official-shaped OOXML fixture includes historical merged header positions.
    buf=io.BytesIO();ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    def sheet(title,area,state='',national=False):
        cells=[('E2',title,'inlineStr'),('E10',area,'inlineStr')]
        if state:cells.append(('F10',state,'inlineStr'))
        for i in range(8):
            col=sources.column_label(7+i*(1 if national else 3))
            total=col if national else sources.column_label(sources.column_number(col)+2)
            serial=(date(2025,1,1)+timedelta(days=90*i)-date(1899,12,30)).days
            cells.extend([(col+'7',str(serial),'n'),(total+'10',str(100+i),'n')])
        return '<worksheet xmlns="'+ns+'"><sheetData>'+''.join('<row r="'+str(n)+'">'+''.join(
            f'<c r="{c}" t="{t}">'+(f'<is><t>{v}</t></is>' if t=='inlineStr' else f'<v>{v}</v>')+'</c>'
            for c,v,t in cells if int(''.join(filter(str.isdigit,c)))==n)+'</row>' for n in (1,2,7,10))+'</sheetData></worksheet>'
    with zipfile.ZipFile(buf,'w') as z:
        z.writestr('xl/workbook.xml',f'<workbook xmlns="{ns}"><workbookPr/></workbook>')
        z.writestr('xl/worksheets/sheet1.xml',sheet('Ladepunkte je Kreis','Kreisfreie Stadt Bonn','Nordrhein-Westfalen'))
        z.writestr('xl/worksheets/sheet2.xml',sheet('Ladepunkte je Bundesland','Nordrhein-Westfalen'))
        z.writestr('xl/worksheets/sheet3.xml',sheet('Ladepunkte in Deutschland','Ladepunkte insgesamt',national=True))
    raw=buf.getvalue()
    rows,region,level=sources.parse_public_chargers(raw,{'city':'Bonn','federal_state':'Nordrhein-Westfalen'})
    assert region=='Bonn' and level=='city' and rows[0]['value']==100
    assert sources.parse_public_chargers(raw,{'city':'Unbekannt','federal_state':'Nordrhein-Westfalen'})[2]=='state'
    assert sources.parse_public_chargers(raw,{})[2]=='country'


def test_context_export_charts_reference_units_only():
    obs={'value':1800,'unit':'Anzahl','year':2025,'page':20}
    ref={'kind':'context','baseline':1113,'target':1207,'unit':'öffentliche Ladepunkte','baseline_date':'2026-01-01',
        'target_date':'2026-07-01','fetched_at':'2026-10-07','assumption':'Private und öffentliche Punkte getrennt.'}
    item={'label':'Ladeinfrastruktur','observation':obs,'reference':ref,'calculation':calculate_reference(obs,ref)}
    prs=Presentation();add_topic(prs,item,'Bonn')
    chart=next(s.chart for s in prs.slides[0].shapes if s.has_chart)
    assert list(chart.series[0].values)==[1113,1207]
    assert '1800' not in str(list(chart.series[0].values))
    assert 'no company projection' in prs.slides[0].notes_slide.notes_text_frame.text


def test_public_chargers_survive_optional_kba_failure(monkeypatch):
    monkeypatch.setattr(references,'public_charger_reference',lambda *a:{'baseline':1,'target':2,'kind':'context'})
    monkeypatch.setattr(references,'vehicle_reference',lambda *a:(_ for _ in ()).throw(ValueError('KBA offline')))
    ref=references.reference_pair('charge_points',2025,2026)
    assert ref['target']==2 and ref['secondary'][0]['error']=='KBA offline'


def test_weather_only_uses_complete_months_and_no_future_days(monkeypatch):
    monkeypatch.setattr(sources,'geocode',lambda location:({'name':'Bonn','latitude':50,'longitude':7},record('{}')))
    days=[date(2024,1,1)+timedelta(days=i) for i in range(31)]+[date(2026,1,1)+timedelta(days=i) for i in range(31)]
    values=[10.0]*len(days);values[-1]=None
    def fetch(url,params):
        assert params['end_date']=='2026-02-06'
        return record({'daily_units':{'temperature_2m_mean':'°C'},'daily':{'time':[d.isoformat() for d in days],'temperature_2m_mean':values}})
    monkeypatch.setattr(sources,'get_record',fetch)
    rows,_=sources.weather_monthly({'city':'Bonn'},2024,2026,date(2026,2,7))
    assert rows==[{'date':'2024-01-01','value':310.0}]


def test_new_topics_require_correct_year_basis_unit_and_defined_scope():
    obs={'metric':'electricity_sales_total','year':2025,'status':'validated','basis':'annual_sales','value':100,
        'unit':'MWh','entity':'Stadtwerke','scope':'group','portfolio':'total'}
    def options(**changes):
        return available_metrics({'metadata':{'year':2025},'observations':[dict(obs,**changes)]})
    assert 'electricity_sales_total' in options()
    for change in ({'year':2024},{'unit':'EUR'},{'basis':'planned'},{'scope':'unknown'},{'portfolio':'project'}):
        assert not options(**change)


def test_credit_reference_includes_uploaded_2023_reporting_year(monkeypatch):
    from stadtwerke import reference_sources as source
    calls=[]
    def record(url,params):
        calls.append(params)
        return {'text':'2023-01;4,12;\n2024-01;5,01;\n','url':url,'params':params,'fetched_at':'2026-10-07T00:00:00Z','sha256':'test'}
    monkeypatch.setattr(source,'get_record',record)
    rows,_=source.credit_rates()
    assert calls[0]['startPeriod']=='2023-01'
    assert rows[0]['date']=='2023-01-01' and rows[0]['value']==4.12


def eurostat_payload(code='CP0441'):
    timeline={r['date'][:7]:i for i,r in enumerate(months())}
    dimensions={key:{'category':{'index':{value:0}}} for key,value in
        [('freq','M'),('unit','I25'),('coicop18',code),('geo','DE')]}
    dimensions['time']={'category':{'index':timeline}}
    return {'id':['freq','unit','coicop18','geo','time'],'size':[1,1,1,1,len(timeline)],
        'dimension':dimensions,'value':{str(i):r['value'] for i,r in enumerate(months())}}


@pytest.mark.parametrize('sewerage',[False,True])
def test_water_prices_use_current_classification_and_complete_historical_series(monkeypatch,sewerage):
    code='CP0443' if sewerage else 'CP0441'
    def fetch(url,params):
        assert url.endswith('/prc_hicp_minr')
        assert params['coicop18']==code and params['unit']=='I25'
        return record(eurostat_payload(code))
    monkeypatch.setattr(sources,'get_record',fetch)
    metric='wastewater_volume' if sewerage else 'water_sales'
    ref=references.reference_pair(metric,2024,2026,today=date(2026,10,7))
    assert ref['baseline']==100 and ref['target']==110
    assert ref['target_date']=='2026-09-30' and ref['months']==9
    assert ref['category']==code and ref['kind']=='context'
    calc=calculate_reference({'value':1000000,'unit':'m³'},ref)
    assert calc['expected'] is None and calc['reference_change_percent']==pytest.approx(10)


def test_water_price_schema_change_is_rejected_and_missing_values_are_not_filled(monkeypatch):
    payload=eurostat_payload()
    del payload['value'][str(payload['dimension']['time']['category']['index']['2026-03'])]
    monkeypatch.setattr(sources,'get_record',lambda *a:record(payload))
    rows,_=sources.water_price_index()
    assert sources.monthly_pair(rows,2024,2026,date(2026,10,7))['months']==2
    payload['dimension']['unit']['category']['index']={'I15':0}
    with pytest.raises(ValueError,match='Indexbasis'):sources.water_price_index()


def test_generated_company_curve_uses_annual_denominator_and_is_exportable():
    from stadtwerke.calculation import company_outlook_series
    pair=sources.monthly_pair(months(),2024,2026,date(2026,10,7))
    ref=dict(pair,baseline=pair['baseline_annual'],aligned_baseline=pair['baseline'],kind='flow',unit='MWh')
    obs={'value':120,'unit':'MWh'}
    calc=calculate_reference(obs,ref);points=calc['company_outlook_series']
    assert points[-1]['value']==pytest.approx(calc['expected'])
    assert points[-1]['historical_proxy']==90
    assert all(p['value_kind']=='conditional_company_projection' for p in points)
    assert company_outlook_series(obs,dict(ref,kind='context'))==[]
    del ref['history'][0]['baseline_value']
    with pytest.raises(ValueError,match='Vergleichsmonat'):company_outlook_series(obs,ref)


@pytest.mark.parametrize('metric',['net_debt_multiple','gearing_ratio','dynamic_debt_ratio'])
def test_financial_ratio_reference_preserves_company_units_without_projection(monkeypatch,metric):
    monkeypatch.setattr(sources,'credit_rates',lambda:(months(),{'provider':'Bundesbank'}))
    monkeypatch.setattr(references,'optional_context',lambda *a:{})
    ref=references.reference_pair(metric,2025,2026,today=date(2026,10,7))
    unit='Faktor' if metric=='net_debt_multiple' else '%'
    calc=calculate_reference({'value':4.9,'unit':unit},ref)
    assert ref['unit']=='% p.a.' and calc['expected'] is None


def test_export_records_generated_chart_origin():
    obs={'value':100,'unit':'m³','year':2025,'page':3}
    ref={'kind':'context','baseline':99.72,'target':104.26,'unit':'Index (2025 = 100)',
         'baseline_date':'2025-08-31','target_date':'2026-08-31','fetched_at':'2026-10-07','provider':'Eurostat'}
    item={'label':'Wasserabsatz / Wasserpreisreferenz','observation':obs,'reference':ref,
          'calculation':calculate_reference(obs,ref)}
    prs=Presentation();add_topic(prs,item,'Bremen')
    texts=' '.join(s.text for s in prs.slides[0].shapes if s.has_text_frame)
    assert 'vom Programm erstellt' in texts and 'API · Referenzvergleich' in texts
    notes=json.loads(prs.slides[0].notes_slide.notes_text_frame.text)
    assert notes['chart_origin']=='generated_by_Stadtwerke_Analyst'
    assert notes['chart_value_kind']=='public_reference_context'


def test_credit_uses_full_year_mean_not_december(monkeypatch):
    rows=[{'date':f'2024-{m:02}-01','value':float(m)} for m in range(1,13)]+[
        {'date':f'2026-{m:02}-01','value':20.0+m} for m in range(1,4)]
    monkeypatch.setattr(sources,'credit_rates',lambda:(rows,{'provider':'Bundesbank'}))
    ref=references.reference_pair('interest_expense',2024,2026,today=date(2026,4,8))
    assert ref['baseline']==6.5
    assert ref['baseline_start']=='2024-01-01' and ref['baseline_date']=='2024-12-31'
    assert ref['target']==22 and ref['aligned_baseline']==2
    assert ref['aggregation']=='mean'
