import csv
import json
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from .config import OUTPUTS

ACCENT=(100,90,155)
POSITIVE=(41,132,89)
NEGATIVE=(166,79,79)

def assessment(item):
    """Indicate a signed model change; context is not a company projection."""
    calc=item['calculation'];ref=item['reference']
    if ref.get('kind')=='context':return None
    change=calc.get('company_change')
    if change is None or change==0:return None
    return {'color':POSITIVE if change>0 else NEGATIVE,
            'label':('Zunahme' if change>0 else 'Abnahme')+' im Modellvergleich',
            'change':('+' if change>0 else '−')+number(abs(change))+' '+item['observation']['unit']}

def number(value, digits=2):
    return f'{value:,.{digits}f}'.replace(',','~').replace('.',',').replace('~','.')

def text(slide,x,y,w,h,content,size=19,bold=False,color=(25,42,58)):
    shape = slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf = shape.text_frame;tf.word_wrap=True
    for i,line in enumerate(content.split('\n')):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold
        p.font.color.rgb=RGBColor(*color)
    return shape

def add_topic(prs, item, company):
    if item['reference'].get('kind')=='scenario':return add_scenario_topic(prs,item,company)
    if item['reference'].get('kind')=='report':return add_personnel_topic(prs,item,company)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    obs=item['observation'];ref=item['reference'];calc=item['calculation']
    text(s,.55,.4,12.2,.6,item['label'].split(' · ')[0],32,True,color=ACCENT)
    entity = obs.get('entity') or company
    subtitle = f"{company}: {entity}" if obs.get('portfolio')=='defined_subset' else entity
    text(s,.55,1.05,12.2,.45,subtitle[:110],17)
    kind=ref.get('kind','stock')
    context=kind=='context'
    chart_unit=ref['unit'] if context else obs['unit']
    text(s,.55,1.48,4.3,.32,'BERICHT · Unternehmensgröße',14,color=ACCENT)
    text(s,5.2,1.48,7.4,.32,('API · Referenzvergleich' if context else 'BERECHNET · Vergleichswert')+f" ({chart_unit})",14,color=ACCENT)
    detail=('Kontextreferenz\nKeine automatische\nUnternehmenshochrechnung' if context else
        ('Zeitraum: Januar bis\n'+ref['target_date'] if kind=='flow' else 'Konstante Relation\n'+number(calc['ratio_percent'],5)+' %'))
    text(s,.55,1.8,4.3,2.7,f"Bericht {obs.get('reporting_period_label',obs['year'])}\n{'über ' if obs.get('qualifier')=='more_than' else 'rund ' if obs.get('approximate') else ''}{number(obs['value'])} {obs['unit']}\n\n{detail}",21)
    data=CategoryChartData()
    if context:
        data.categories=[ref['baseline_date'],ref['target_date']]
        values=[ref['baseline'],ref['target']]
    elif kind=='flow':
        data.categories=[f"Jan–{ref['months']:02} / {obs['year']}",f"Jan–{ref['months']:02} / {ref['target_date'][:4]}"]
        values=[calc['historical_company_aligned'],calc['expected']]
    else:
        data.categories=[str(obs['year']),ref['target_date']]
        values=[obs['value'],calc['expected']]
    data.add_series('Referenz' if context else 'Bedingter Vergleich',values)
    chart=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(5.2),Inches(1.8),Inches(7.4),Inches(3.3),data).chart
    chart.has_legend=False
    chart.has_title=False
    chart.category_axis.tick_labels.font.size=Pt(16)
    chart.value_axis.tick_labels.font.size=Pt(14)
    if min(values)>=0: chart.value_axis.minimum_scale=0
    chart.plots[0].has_data_labels=True
    chart.plots[0].data_labels.font.size=Pt(17)
    chart.plots[0].data_labels.number_format='0.00'
    for index,point in enumerate(chart.series[0].points):
        point.format.fill.solid();point.format.fill.fore_color.rgb=RGBColor(*( (100,90,155) if index==1 else (216,213,225)))
    from .calculation import reference_comparison
    comparison=reference_comparison(ref)
    summary=(f"API · {comparison['baseline_period']} / {comparison['target_period']}: {number(comparison['baseline'])} → {number(comparison['target'])} {ref['unit']}. "
        +('Keine abgeleitete Unternehmenszahl.' if context else f"Modellwert: {number(calc['expected'])} {obs['unit']}.")
        +(' Gleiche Monate; keine Jahresendprognose.' if kind=='flow' else ''))
    if kind=='flow':
        summary=f"API · {comparison['baseline_period']} / {comparison['target_period']}:\n{number(comparison['baseline'])} → {number(comparison['target'])} {ref['unit']} · gleiche Monate"
    text(s,.55,5.2,12.2,.65,summary,17)
    evaluation=assessment(item)
    if evaluation:text(s,.55,4.65,4.3,.5,evaluation['label']+': '+evaluation['change'],12,color=evaluation['color'])
    if context:
        change=calc['reference_change'];sign='+' if change>0 else ''
        text(s,5.2,5.0,7.4,.35,'BERECHNET · Änderung: '+sign+number(change)+' '+('Prozentpunkte' if ref['unit']=='% p.a.' else ref['unit']),14,color=ACCENT)
    assumption=ref.get('assumption','Konstante Relation zur nationalen Referenz.')
    if kind=='flow':assumption='Jahresreferenz als Berechnungsbasis: '+number(ref['baseline'])+' '+ref['unit']+'. '+assumption
    text(s,.55,5.92,12.2,.72,assumption,12)
    text(s,.55,6.7,12.2,.55,
        f"Grafik und Analyse: vom Programm erstellt. Daten: Geschäftsbericht, PDF-Seite {obs['page']}; {ref.get('provider','Energy-Charts')}. "
        f"Referenzen {ref['baseline_date']} / {ref['target_date']}; Abruf {ref['fetched_at'][:10]}. Quellen und Originalbelege in Foliennotizen.",11)
    s.notes_slide.notes_text_frame.text = json.dumps({'chart_origin':'generated_by_Stadtwerke_Analyst',
        'chart_value_kind':'public_reference_context' if context else 'conditional_company_projection',
        'report_evidence':obs,'reference':ref,'formula':calc['formula']},ensure_ascii=False,indent=2)

def export(result, session_id):
    from .outlook import company_result
    original=result
    result=company_result(result)
    directory=OUTPUTS/session_id;directory.mkdir(parents=True,exist_ok=True)
    for item in original.get('topics',[]):
        if item['reference'].get('kind') in ('context','scenario'):(directory/(item['metric']+'.pptx')).unlink(missing_ok=True)
    if not result['topics']:(directory/'Stadtwerke_Ausblick.pptx').unlink(missing_ok=True)
    json_path=directory/'analyse.json';json_path.write_text(json.dumps(result,ensure_ascii=False,indent=2))
    csv_path=directory/'vergleichswerte.csv'
    with csv_path.open('w',newline='') as stream:
        writer=csv.writer(stream,delimiter=';');writer.writerow(['Thema','Berichtseinheit','Berichtswert','Bedingter Vergleichswert','Datenstand','Vergleichsart','Referenzeinheit','Historische Referenz','Aktuelle Referenz','Historischer Referenzzeitraum','Berechnete Kennzahl','Wert','Einheit','Mitarbeiter im Bericht','Auszubildende im Bericht','Szenarioeinheit','Berechnete eigene Basis','Unternehmensänderung','Eigene Annahmen','Aktueller Referenzzeitraum','Jahresreferenz als Berechnungsbasis'])
        from .calculation import reference_comparison
        for item in result['topics']:
            comparison=reference_comparison(item['reference'])
            c=item['calculation'];writer.writerow([item['label'],item['observation']['unit'],item['observation']['value'],c['expected'],item['reference']['target_date'],item['reference'].get('kind','stock'),item['reference']['unit'],comparison['baseline'],comparison['target'],comparison['baseline_period'],'Ausbildungsquote' if item['reference']['kind']=='report' else '',c.get('apprentice_share_percent') if item['reference']['kind']=='report' else '','%' if c.get('apprentice_share_percent') is not None else '',(item['reference'].get('personnel',{}).get('employees') or {}).get('value',''),(item['reference'].get('personnel',{}).get('apprentices') or {}).get('value',''),c.get('output_unit',item['observation']['unit']),c.get('baseline_company',''),c.get('company_change',''),json.dumps(c.get('assumptions',{}),ensure_ascii=False),comparison['target_period'],item['reference']['baseline'] if item['reference'].get('kind')=='flow' else ''])
    files={'json':str(json_path),'csv':str(csv_path)}
    if result['topics']:
        prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
        for item in result['topics']:
            add_topic(prs,item,result['company'])
            single=Presentation();single.slide_width=prs.slide_width;single.slide_height=prs.slide_height
            add_topic(single,item,result['company']);single.save(directory/(item['metric']+'.pptx'))
        pptx=directory/'Stadtwerke_Ausblick.pptx';prs.save(pptx);files['pptx']=str(pptx)
    return files


def add_personnel_topic(prs,item,company):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    ref=item['reference'];calc=item['calculation'];rows=ref['personnel']
    text(slide,.55,.4,12.2,.65,'Beschäftigte & Ausbildung',32,True,color=ACCENT)
    text(slide,.55,1.1,12.2,.5,company+' · Bericht '+ref['target_date'],18)
    for index,(key,label) in enumerate((('employees','Mitarbeiter'),('apprentices','Auszubildende'))):
        row=rows.get(key)
        if row:
            text(slide,.55+index*4.3,2,4.1,.5,'BERICHT · '+label,14)
            text(slide,.55+index*4.3,2.6,4.1,.7,number(row['value'],0)+' Personen',28)
            text(slide,.55+index*4.3,3.4,4.1,.4,'Original-PDF · Seite '+str(row['page']),14)
    share=calc.get('apprentice_share_percent')
    text(slide,9.2,2,3.4,.5,'BERECHNET · Ausbildungsquote',14,color=ACCENT)
    text(slide,9.2,2.6,3.4,.7,number(share)+' %' if share is not None else 'Abgrenzung offen',28,color=ACCENT)
    text(slide,.55,4.6,12.2,1.1,ref['assumption'],18)
    text(slide,.55,6.1,12.2,.8,'Originalzahlen aus dem Geschäftsbericht; Quote vom Programm berechnet. Zeitbasis und Abgrenzung bleiben im Beleg erhalten.',14)
    slide.notes_slide.notes_text_frame.text=json.dumps({'personnel_evidence':rows,'calculation':calc},ensure_ascii=False,indent=2)


def add_scenario_topic(prs,item,company):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    obs,ref,c=item['observation'],item['reference'],item['calculation']
    text(slide,.55,.4,12.2,.65,item['label'].replace(' · Angaben erforderlich',''),28,True,color=ACCENT)
    text(slide,.55,1.05,12.2,.5,company+' · '+c['period_label'],17)
    text(slide,.55,1.7,4,.35,'BERICHT · Unternehmensgröße',14,color=ACCENT)
    text(slide,.55,2.1,4,.8,number(obs['value'])+' '+obs['unit'],24)
    text(slide,.55,3,4,1,'API · '+ref.get('provider','Referenz')+'\n'+number(ref['baseline'])+' → '+number(ref['target'])+' '+ref['unit'],16)
    text(slide,.55,4.15,4,1.2,'EIGENE ANNAHME\nBetroffene Menge: '+number(c['assumptions']['amount'])+' '+obs['unit'],16)
    data=CategoryChartData();data.categories=['Eigene Basis · berechnet','Szenario · berechnet']
    data.add_series(c['output_unit'],[c['baseline_company'],c['expected']])
    chart=slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(5),Inches(2),Inches(7.7),Inches(3.1),data).chart
    chart.has_legend=False;chart.has_title=False
    for index,point in enumerate(chart.series[0].points):
        point.format.fill.solid();point.format.fill.fore_color.rgb=RGBColor(*(ACCENT if index else (216,213,225)))
    text(slide,5,1.65,7.7,.35,'BERECHNET · '+c['title']+' ('+c['output_unit']+')',14,color=ACCENT)
    text(slide,.55,5.45,12.2,.7,'Szenario: '+number(c['expected'])+' '+c['output_unit']+' · Änderung: '+number(c['company_change'])+' '+c['output_unit'],20,color=ACCENT)
    text(slide,.55,6.15,12.2,.85,'Bedingte Unternehmensrechnung mit bestätigten eigenen Eingaben. PDF-Seite '+str(obs['page'])+'; API '+ref['baseline_date']+' / '+ref['target_date']+'. Vollständige Annahmen, Originalbelege und Formel in den Foliennotizen.',13)
    slide.notes_slide.notes_text_frame.text=json.dumps({'report_evidence':obs,'reference':ref,'calculation':c,'input_origin':'user_assumption'},ensure_ascii=False,indent=2)
